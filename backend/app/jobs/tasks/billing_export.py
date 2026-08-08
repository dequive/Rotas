"""ARQ tasks for billing compliance exports.

Task: task_export_compliance_report — generates monthly AT compliance XLSX.
"""

from __future__ import annotations

import logging
from datetime import datetime
from io import BytesIO
from typing import Any
from uuid import UUID

from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import select, text

from app.modules.billing.models import BillingDocument, ExportJob
from app.modules.contracts.models import Contract
from app.modules.files.service import save_generated_file

logger = logging.getLogger(__name__)

COMPLIANCE_REPORT_COLUMNS = [
    "invoice_number",
    "client_nuit",
    "client_name",
    "issued_at",
    "subtotal",
    "iva_rate",
    "iva_amount",
    "total_amount",
    "status",
]


async def task_export_compliance_report(
    ctx: dict[str, Any],
    job_id: str,
    month: str,
    tenant_id: str,
) -> dict:
    """Generate monthly compliance XLSX for AT submission.

    Args:
        job_id: ExportJob UUID string
        month: YYYY-MM format (e.g. "2026-01")
        tenant_id: Tenant UUID string
    """
    async with ctx["session_factory"]() as db:
        job = await db.scalar(select(ExportJob).where(ExportJob.id == UUID(job_id)))
        if not job:
            logger.error("ExportJob %s not found", job_id)
            return {"status": "failed", "error": "job_not_found"}
        if job.status == "done":
            return {
                "status": "done",
                "file_id": str(job.file_id) if job.file_id else None,
            }

        job.status = "processing"
        await db.commit()

        try:
            month_dt = datetime.strptime(month, "%Y-%m")
            stmt = (
                select(BillingDocument, Contract.client_nuit)
                .outerjoin(Contract, BillingDocument.contract_id == Contract.id)
                .where(
                    BillingDocument.tenant_id == UUID(tenant_id),
                    BillingDocument.status != "draft",
                    text(
                        "DATE_TRUNC('month', billing_documents.issued_at) = :month_start"
                    ).bindparams(month_start=month_dt),
                )
                .order_by(BillingDocument.issued_at)
            )
            result = await db.execute(stmt)
            rows = result.all()

            wb = Workbook()
            ws = wb.active
            if ws is None:
                ws = wb.create_sheet()
            ws.title = f"Compliance {month}"

            for col_idx, col_name in enumerate(COMPLIANCE_REPORT_COLUMNS, start=1):
                cell = ws.cell(row=1, column=col_idx, value=col_name.upper())
                cell.font = Font(bold=True)

            for row_idx, (doc, client_nuit) in enumerate(rows, start=2):
                ws.cell(row=row_idx, column=1, value=doc.invoice_number or "")
                ws.cell(row=row_idx, column=2, value=client_nuit or "")
                ws.cell(row=row_idx, column=3, value=doc.client_name or "")
                ws.cell(
                    row=row_idx,
                    column=4,
                    value=doc.issued_at.isoformat() if doc.issued_at else "",
                )
                ws.cell(row=row_idx, column=5, value=float(doc.subtotal or 0))
                ws.cell(
                    row=row_idx,
                    column=6,
                    value=float(doc.iva_rate) if doc.iva_rate is not None else "",
                )
                ws.cell(row=row_idx, column=7, value=float(doc.tax_amount or 0))
                ws.cell(row=row_idx, column=8, value=float(doc.total_amount or 0))
                ws.cell(row=row_idx, column=9, value=doc.status)

            buffer = BytesIO()
            wb.save(buffer)
            buffer.seek(0)
            xlsx_bytes = buffer.read()

            filename = f"compliance_{month}_{tenant_id[:8]}.xlsx"
            file_obj = await save_generated_file(
                db,
                UUID(tenant_id),
                content=xlsx_bytes,
                filename=filename,
                mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                file_type="compliance_report",
                entity_type="compliance_report",
            )

            job.status = "done"
            job.file_id = file_obj.id
            await db.commit()
            logger.info("Compliance report %s done: %d rows", job_id, len(rows))
            return {"status": "done", "file_id": str(file_obj.id), "row_count": len(rows)}

        except Exception as exc:
            logger.exception("Compliance report %s failed", job_id)
            job.status = "failed"
            job.error_message = str(exc)[:500]
            await db.commit()
            return {"status": "failed", "error": str(exc)[:200]}
